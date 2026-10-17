from types import SimpleNamespace

from openpyxl import Workbook

from app.api.batches import _normalize_ocr_result
from app.services.export_excel import _add_container_rows, _populate_container_sheet


def test_bundle_register_keeps_valid_lot_numbers_and_rejects_names() -> None:
    result = _normalize_ocr_result(
        {
            "container_no": "TEMU1234567",
            "bundles": [
                {"batch_no": "r2-54", "bundle_no": "1", "ingot_count": 30},
                {"batch_no": "RICH", "bundle_no": "2", "ingot_count": 30},
            ],
        }
    )

    assert result["container_no"] == "TEMU1234567"
    assert result["bundles"][0]["batch_no"] == "R2-54"
    assert result["bundles"][1]["batch_no"] is None


def test_container_detail_writes_each_bundle_lot_and_expands_rows() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "C-1&2"
    worksheet.merge_cells("B5:B23")
    worksheet.merge_cells("H5:H23")
    bundles = [
        SimpleNamespace(batch_no="R2-54", bundle_no=str(number), ingot_count=30, gross_weight=1200.0)
        for number in range(1, 22)
    ]
    bundles[1].batch_no = "R2-55"

    _add_container_rows(worksheet, len(bundles))
    _populate_container_sheet(worksheet, [(1, "DEMO-001", bundles)])

    assert worksheet["A2"].value == "Container No.1 : DEMO-001"
    assert worksheet["B5"].value == "R2-54"
    assert worksheet["B6"].value == "R2-55"
    assert worksheet["C25"].value == "21"
    assert worksheet["A26"].value == "Total"
    assert worksheet["D26"].value == "=SUM(D5:D25)"
