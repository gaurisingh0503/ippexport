import io
from collections.abc import Iterable
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.cell_range import CellRange

from app.models.export_document import ExportDocument
from app.models.sticker_batch import StickerBatch
from app.services.storage import download_storage_object

_INVOICE_PACKING_TEMPLATE = "templates/invoice_packing_template.xlsx"
_COA_TEMPLATE = "templates/coa_only_template.xlsx"
_BUNDLES_PER_CONTAINER_SECTION = 20


def _sheet(workbook: Workbook, title: str) -> Worksheet:
    try:
        return workbook[title]
    except KeyError as exc:
        raise ValueError(f"Template is missing the {title} sheet") from exc


def _write_lines(ws: Worksheet, column: str, start_row: int, lines: Iterable[str]) -> None:
    for offset, line in enumerate(lines):
        ws[f"{column}{start_row + offset}"] = line


def _write_party(ws: Worksheet, column: str, start_row: int, party: dict[str, list[str] | str]) -> None:
    cell = ws[f"{column}{start_row}"]
    lines = [party["name"], *party["lines"]]
    cell.value = "\n".join(lines)
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _merge_value(ws: Worksheet, cell_range: str, value: object) -> None:
    start_cell = cell_range.split(":")[0]
    target_range = CellRange(cell_range)
    for merged_range in list(ws.merged_cells.ranges):
        overlaps = not (
            target_range.max_col < merged_range.min_col
            or target_range.min_col > merged_range.max_col
            or target_range.max_row < merged_range.min_row
            or target_range.min_row > merged_range.max_row
        )
        if overlaps and str(merged_range) != cell_range:
            ws.unmerge_cells(str(merged_range))
    if cell_range not in {str(merged_range) for merged_range in ws.merged_cells.ranges}:
        ws.merge_cells(cell_range)
    cell = ws[start_cell]
    cell.value = value
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _amount_in_words(amount: float) -> str:
    ones = [
        "", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
        "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
        "Seventeen", "Eighteen", "Nineteen",
    ]
    tens = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]

    def _below_thousand(n: int) -> str:
        if n == 0:
            return ""
        if n < 20:
            return ones[n]
        if n < 100:
            return tens[n // 10] + (" " + ones[n % 10] if n % 10 else "")
        rest = _below_thousand(n % 100)
        return ones[n // 100] + " Hundred" + (" " + rest if rest else "")

    cents = round((amount % 1) * 100)
    dollars = int(amount)
    if dollars == 0:
        words = "Zero"
    else:
        parts = []
        if dollars // 1_000_000_000:
            parts.append(_below_thousand(dollars // 1_000_000_000) + " Billion")
        if (dollars % 1_000_000_000) // 1_000_000:
            parts.append(_below_thousand((dollars % 1_000_000_000) // 1_000_000) + " Million")
        if (dollars % 1_000_000) // 1_000:
            parts.append(_below_thousand((dollars % 1_000_000) // 1_000) + " Thousand")
        if dollars % 1_000:
            parts.append(_below_thousand(dollars % 1_000))
        words = " ".join(parts)
    return f"US Dollars {words} and {cents:02d}/100 Only"


def _batches(document: ExportDocument) -> list[StickerBatch]:
    return [item.sticker_batch for item in sorted(document.batches, key=lambda item: item.position)]


def _bundles(document: ExportDocument):
    return [bundle for batch in _batches(document) for bundle in batch.bundles]


def _populate_address_fields(ws: Worksheet, document: ExportDocument) -> None:
    shipper = document.shipper
    ws["A4"] = shipper.name
    _write_lines(ws, "A", 5, [shipper.address_line1, shipper.address_line2, shipper.email])
    ws["F4"] = document.invoice_no
    _merge_value(ws, "H4:I5", document.invoice_date)
    _merge_value(ws, "A9:E10", None)
    _write_party(ws, "A", 9, document.consignee)
    if document.notify1:
        _write_party(ws, "A", 12, document.notify1)
    if document.notify2:
        _write_party(ws, "A", 20, document.notify2)
    ws["F10"] = document.port_of_discharge
    ws["F12"] = document.term_of_delivery
    ws["F12"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if document.term_of_delivery:
        lines = max(1, -(-len(document.term_of_delivery) // 25))  # ceil div, ~25 chars per line
        ws.row_dimensions[12].height = max(15, lines * 15)
    _merge_value(ws, "H12:I13", document.bill_of_lading_no)
    _merge_value(ws, "F15:I24", document.contract_no)


def _populate_container_sheet(
    ws: Worksheet, container_groups: list[tuple[int, str, list]]
) -> None:
    slots = (("A", "B", "C", "D", "E"), ("G", "H", "I", "J", "K"))
    for index, (container_index, container_no, bundles) in enumerate(container_groups):
        if index >= len(slots):
            return
        serial_col, batch_col, bundle_col, ingot_col, weight_col = slots[index]
        ws[f"{serial_col}2"] = f"Container No.{container_index} : {container_no}"
        for row_offset, bundle in enumerate(bundles):
            row = 5 + row_offset
            ws[f"{serial_col}{row}"] = row_offset + 1
            ws[f"{batch_col}{row}"] = bundle.batch_no
            ws[f"{bundle_col}{row}"] = bundle.bundle_no
            ws[f"{ingot_col}{row}"] = bundle.ingot_count
            ws[f"{weight_col}{row}"] = bundle.gross_weight


def _add_container_rows(ws: Worksheet, required_rows: int) -> None:
    additional_rows = max(0, required_rows - _BUNDLES_PER_CONTAINER_SECTION)
    for merged_range in ("B5:B23", "H5:H23"):
        ws.unmerge_cells(merged_range)
    if additional_rows:
        ws.insert_rows(24, additional_rows)

        for row in range(24, 24 + additional_rows):
            for column in range(1, 12):
                source = ws.cell(22, column)
                target = ws.cell(row, column)
                if source.has_style:
                    target._style = copy(source._style)
                if source.number_format:
                    target.number_format = source.number_format
                if source.alignment:
                    target.alignment = copy(source.alignment)

    last_row = max(_BUNDLES_PER_CONTAINER_SECTION + 4, required_rows + 4)
    for row in range(5, last_row + 1):
        ws[f"B{row}"]._style = copy(ws[f"C{row}"]._style)
        ws[f"H{row}"]._style = copy(ws[f"I{row}"]._style)

    total_row = last_row + 1
    ws[f"A{total_row}"] = "Total"
    ws[f"D{total_row}"] = f"=SUM(D5:D{last_row})"
    ws[f"E{total_row}"] = f"=SUM(E5:E{last_row})"
    ws[f"G{total_row}"] = "Total"
    ws[f"J{total_row}"] = f"=SUM(J5:J{last_row})"
    ws[f"K{total_row}"] = f"=SUM(K5:K{last_row})"


def _populate_invoice_packing(workbook: Workbook, document: ExportDocument) -> None:
    invoice = _sheet(workbook, "Provisional Inovice")
    packing_list = _sheet(workbook, "Packing List")
    container_12 = _sheet(workbook, "C-1&2")
    container_34 = _sheet(workbook, "C-3&4")
    bundles = _bundles(document)
    batches = _batches(document)
    material = bundles[0].material
    if material is None:
        raise ValueError("Every bundle needs a material before generating documents")

    _populate_address_fields(invoice, document)
    _populate_address_fields(packing_list, document)
    goods_description = document.customer_material_name or material.goods_description or material.material_name
    invoice["B26"] = goods_description
    invoice["F26"] = document.unit_price_usd_t
    invoice["G26"] = sum(bundle.net_weight or 0 for bundle in bundles) / 1000
    invoice["I26"] = document.unit_price_usd_t * invoice["G26"].value
    invoice["I27"] = invoice["I26"].value
    invoice["A29"] = _amount_in_words(invoice["I27"].value or 0)
    invoice["D34"] = invoice["G26"].value
    invoice["D35"] = sum(bundle.gross_weight or 0 for bundle in bundles) / 1000
    invoice["D36"] = len(bundles)
    invoice["D37"] = document.payment_terms
    invoice["D38"] = f"{document.shipper.name}, Thailand"
    invoice["D39"] = document.shipping_mark
    invoice["D40"] = material.hs_code
    packing_list["D26"] = invoice["D34"].value
    packing_list["D27"] = invoice["D35"].value
    packing_list["D28"] = len(bundles)
    packing_list["D29"] = invoice["D38"].value
    packing_list["D30"] = document.shipping_mark
    packing_list["D31"] = material.hs_code

    container_groups: list[tuple[int, str, list]] = []
    for container_index, batch in enumerate(batches, start=1):
        if not batch.container_no:
            raise ValueError("Every grouped batch needs a container number")
        container_groups.append((container_index, batch.container_no, batch.bundles))

    container_sheets = [container_12, container_34]
    required_sheets = max(1, (len(container_groups) + 1) // 2)
    while len(container_sheets) < required_sheets:
        clone = workbook.copy_worksheet(container_34)
        clone.title = f"C-{len(container_sheets) * 2 + 1}&{len(container_sheets) * 2 + 2}"
        container_sheets.append(clone)
    for worksheet in container_sheets[required_sheets:]:
        workbook.remove(worksheet)
    for page_index, worksheet in enumerate(container_sheets[:required_sheets]):
        page_groups = container_groups[page_index * 2 : page_index * 2 + 2]
        _add_container_rows(worksheet, max(len(group[2]) for group in page_groups))
        _populate_container_sheet(worksheet, page_groups)


def _populate_coa(workbook: Workbook, document: ExportDocument) -> None:
    batches = _batches(document)
    coa_sheets = [worksheet for worksheet in workbook.worksheets if worksheet.title.startswith("COA_")]
    if not coa_sheets:
        raise ValueError("Template is missing a COA sheet")

    while len(coa_sheets) < len(batches):
        coa_sheets.append(workbook.copy_worksheet(coa_sheets[-1]))

    for worksheet, batch in zip(coa_sheets, batches, strict=False):
        bundle = batch.bundles[0]
        material = bundle.material
        worksheet.title = f"COA_{bundle.batch_no}"
        for merged_range in list(worksheet.merged_cells.ranges):
            worksheet.unmerge_cells(str(merged_range))
        _merge_value(worksheet, "A4:E4", document.consignee["name"])
        _merge_value(worksheet, "F4:G4", document.invoice_no)
        _merge_value(worksheet, "H4:I4", document.invoice_date)
        _merge_value(worksheet, "F6:G11", document.production_date)
        _merge_value(worksheet, "H6:I11", bundle.batch_no)
        if document.notify1:
            lines = [document.notify1["name"], *document.notify1["lines"]]
            for i, line in enumerate(lines[:6]):
                worksheet[f"A{6 + i}"] = line
        if document.notify2:
            lines = [document.notify2["name"], *document.notify2["lines"]]
            for i, line in enumerate(lines[:6]):
                worksheet[f"A{13 + i}"] = line
        worksheet["A20"] = (
            f"Goods Description - {document.customer_material_name or material.goods_description or material.material_name}"
            if material
            else "Goods Description"
        )
        for row in range(22, 35):
            symbol = worksheet[f"E{row}"].value
            if symbol and document.coa_specifications:
                worksheet[f"F{row}"] = document.coa_specifications.get(symbol, worksheet[f"F{row}"].value)
            if symbol and document.coa_actuals:
                worksheet[f"H{row}"] = document.coa_actuals.get(bundle.batch_no, {}).get(symbol)

    for worksheet in coa_sheets[len(batches) :]:
        workbook.remove(worksheet)


async def generate_invoice_packing(document: ExportDocument) -> bytes:
    template = await download_storage_object(_INVOICE_PACKING_TEMPLATE)
    workbook = load_workbook(io.BytesIO(template))
    _populate_invoice_packing(workbook, document)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def generate_coa(document: ExportDocument) -> bytes:
    template = await download_storage_object(_COA_TEMPLATE)
    workbook = load_workbook(io.BytesIO(template))
    _populate_coa(workbook, document)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
