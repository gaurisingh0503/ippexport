import io
import re
from collections.abc import Iterable
from copy import copy

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.models.bundle import Bundle
from app.models.sticker_batch import StickerBatch
from app.services.storage import download_storage_object

_TEMPLATE_KEYS = {
    ("domestic", "ipp"): "templates/domestic_sticker_ipp.xlsx",
    ("domestic", "richsiri"): "templates/domestic_sticker_richsiri.xlsx",
    ("export", "ipp"): "templates/sticker_template.xlsx",
    ("export", "richsiri"): "templates/sticker_template.xlsx",
}

_FIELD_LABELS = {
    "material": "material name",
    "batch_no": "batch no",
    "bundle_no": "bundle no",
    "ingot_count": "number of ingots",
    "gross_weight": "gross weight",
    "net_weight": "net weight",
}


def template_key_for_batch(batch: StickerBatch) -> str:
    key = (batch.type, batch.shipper.slug)
    try:
        return _TEMPLATE_KEYS[key]
    except KeyError as exc:
        raise ValueError(f"No sticker template for {batch.type} batch / {batch.shipper.slug}") from exc


def _normalize_label(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def _label_cells(ws: Worksheet, label: str) -> list[tuple[int, int]]:
    return [
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if _normalize_label(cell.value) == label
    ]


def _group_end_column(ws: Worksheet, row: int, start_column: int) -> int:
    next_starts = [
        column
        for _, column in _label_cells(ws, _FIELD_LABELS["material"])
        if _ == row and column > start_column
    ]
    if next_starts:
        return next_starts[0] - 1

    column = start_column + 1
    while column <= ws.max_column and ws.cell(row, column).border.right.style is None:
        column += 1
    return column


def _slot_targets(ws: Worksheet) -> list[dict[str, str]]:
    material_cells = _label_cells(ws, _FIELD_LABELS["material"])
    slots: list[dict[str, str]] = []

    for material_row, material_column in material_cells:
        group_end = _group_end_column(ws, material_row, material_column)
        target_columns = {
            "material": material_column + 2,
            "batch_no": group_end - 1,
            "bundle_no": group_end - 1,
            "ingot_count": group_end - 1,
            "gross_weight": group_end - 1,
            "net_weight": group_end - 1,
        }

        targets: dict[str, str] = {}
        for field, label in _FIELD_LABELS.items():
            candidates = [
                row
                for row, column in _label_cells(ws, label)
                if column == material_column and row >= material_row
            ]
            if candidates:
                targets[field] = ws.cell(min(candidates), target_columns[field]).coordinate

        if len(targets) == len(_FIELD_LABELS):
            slots.append(targets)

    return slots


def _image_from_bytes(data: bytes, source: Image) -> Image:
    buffer = io.BytesIO(data)
    image = Image(buffer)
    image.width = source.width
    image.height = source.height
    image.anchor = copy(source.anchor)
    image._source_buffer = buffer
    return image


def _copy_images(source: Worksheet, destination: Worksheet) -> None:
    rebuilt_source_images = []
    destination_images = []
    for image in source._images:
        data = image._data()
        rebuilt_source_images.append(_image_from_bytes(data, image))
        destination_images.append(_image_from_bytes(data, image))

    source._images = rebuilt_source_images
    for image in destination_images:
        destination.add_image(image)


def _add_sheet_clone(workbook: Workbook, source: Worksheet) -> Worksheet:
    clone = workbook.copy_worksheet(source)
    _copy_images(source, clone)
    return clone


def _ensure_slots(workbook: Workbook, required: int) -> list[tuple[Worksheet, dict[str, str]]]:
    slots = [
        (worksheet, target)
        for worksheet in workbook.worksheets
        for target in _slot_targets(worksheet)
    ]
    source_sheets = list(workbook.worksheets)
    source_index = 0

    while len(slots) < required:
        source = source_sheets[source_index % len(source_sheets)]
        clone = _add_sheet_clone(workbook, source)
        slots.extend((clone, target) for target in _slot_targets(clone))
        source_index += 1

    return slots


def _write_bundle(slot: tuple[Worksheet, dict[str, str]], bundle: Bundle) -> None:
    worksheet, targets = slot
    material = bundle.material
    if material is None:
        raise ValueError(f"Bundle {bundle.id} has no material")

    values: dict[str, object] = {
        "material": material.sticker_display_name,
        "batch_no": bundle.batch_no,
        "bundle_no": bundle.bundle_no,
        "ingot_count": bundle.ingot_count,
        "gross_weight": bundle.gross_weight,
        "net_weight": bundle.net_weight,
    }
    for field, target in targets.items():
        worksheet[target] = values[field]
        if field == "bundle_no":
            alignment = copy(worksheet[target].alignment)
            alignment.horizontal = "right"
            worksheet[target].alignment = alignment


async def generate_stickers(batch: StickerBatch, bundles: Iterable[Bundle]) -> bytes:
    bundle_list = list(bundles)
    if not bundle_list:
        raise ValueError("Add at least one bundle before generating stickers")

    template = await download_storage_object(template_key_for_batch(batch))
    workbook = load_workbook(io.BytesIO(template))
    slots = _ensure_slots(workbook, len(bundle_list))

    for worksheet, targets in slots:
        for target in targets.values():
            worksheet[target] = None

    for slot, bundle in zip(slots[: len(bundle_list)], bundle_list, strict=True):
        _write_bundle(slot, bundle)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
